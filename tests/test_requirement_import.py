"""The source workbook is read fully; the checked-in import stays assigned-only."""

import hashlib
import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from model.tools.import_requirements import SHEETS, import_workbook


class RequirementImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.workbook = self.folder / "Verification.xlsx"
        self.scope = self.folder / "verification.toml"
        self.destination = self.folder / "requirements.json"
        self.scope.write_text('''description = "One assigned contribution"
[[activities]]
activity_id = "TEST.1.01.A01"
calculation = "thrust_sum"
expected_activity = "Installed thrust calculation."
artifact = "Installed thrust"
coverage = "Nominal arithmetic only."
''', encoding="utf-8")

    def write_workbook(self, duplicate_unassigned=False, stale=False):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        headers = ["ID", "Requirement", "Verification Activity", "Verification Method", "Verification Artifact"]
        for name in SHEETS:
            sheet = workbook.create_sheet(name)
            for column, header in enumerate(headers, 1):
                sheet.cell(4, column, header)
        first = workbook[SHEETS[0]]
        first.append(["TEST.1.01", "Selected requirement", "Changed calculation." if stale else "Installed thrust calculation.\nPhysical inspection.", "Analysis / Inspection", "Calculation and record"])
        first.append(["TEST.1.02", "Unassigned requirement TBD-LIMIT", "Unassigned analysis.", "Analysis", "Unassigned output"])
        workbook[SHEETS[-1]].append(["TEST.1.02" if duplicate_unassigned else "TEST.1.03", "Another unassigned requirement", "Unassigned test.", "Test", "Test report"])
        workbook.save(self.workbook)
        workbook.close()

    def test_writes_selected_requirements_with_all_source_activities_and_provenance(self):
        self.write_workbook()
        summary = import_workbook(self.workbook, self.destination, self.scope)
        records = json.loads(self.destination.read_text(encoding="utf-8"))
        self.assertEqual(summary["workbook_requirements_read"], 3)
        self.assertEqual(summary["count"], 1)
        self.assertEqual(summary["assigned_activities"], 1)
        self.assertEqual([row["id"] for row in records], ["TEST.1.01"])
        self.assertEqual(records[0]["activities"], ["Installed thrust calculation.", "Physical inspection."])
        self.assertEqual(records[0]["source_ref"], "Tier 0!C5")
        self.assertEqual(records[0]["source_file"], self.workbook.name)
        self.assertEqual(records[0]["source_sha256"], hashlib.sha256(self.workbook.read_bytes()).hexdigest())

    def test_duplicate_even_outside_assignment_rejects_entire_import(self):
        self.write_workbook(duplicate_unassigned=True)
        with self.assertRaisesRegex(ValueError, "duplicated"):
            import_workbook(self.workbook, self.destination, self.scope)
        self.assertFalse(self.destination.exists())

    def test_stale_assignment_does_not_overwrite_previous_import(self):
        self.write_workbook(stale=True)
        self.destination.write_text("previous import", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "wording mismatch"):
            import_workbook(self.workbook, self.destination, self.scope)
        self.assertEqual(self.destination.read_text(encoding="utf-8"), "previous import")


if __name__ == "__main__":
    unittest.main()
