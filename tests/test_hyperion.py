"""Checks for the visible calculations in Hyperion.ipynb.

Run from the repository root: python -m unittest discover -s tests.
Only cells tagged ``definitions`` are loaded here. Tests do not run the complete
flight, display plots, or export project results. Run All in Jupyter checks the
complete flight and records its visible outputs.
"""

import ast
from contextlib import contextmanager
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
import tempfile
from types import ModuleType
import unittest
from unittest.mock import patch

import numpy as np
import openpyxl
import pandas as pd
from scipy.integrate import solve_ivp
from scipy.spatial.transform import Rotation

ROOT = Path(__file__).resolve().parents[1]
NOTEBOOK_PATH = ROOT / "Hyperion.ipynb"


def load_notebook_definitions(path=NOTEBOOK_PATH):
    """Load visible definition cells into a patchable module for isolated tests."""
    document = json.loads(path.read_text(encoding="utf-8"))
    module = ModuleType("hyperion_notebook")
    module.__file__ = str(path)
    sys.modules[module.__name__] = module
    previous_directory = Path.cwd()
    try:
        os.chdir(path.parent)
        for number, cell in enumerate(document["cells"], 1):
            tags = cell.get("metadata", {}).get("tags", [])
            if cell["cell_type"] == "code" and "definitions" in tags:
                source = "".join(cell["source"])
                filename = f"{path}:cell-{number}"
                exec(compile(source, filename, "exec"), module.__dict__)
    finally:
        os.chdir(previous_directory)
    return document, module


DOCUMENT, NOTEBOOK = load_notebook_definitions()


class SharedInputs(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.inputs = Path(self.temporary.name)
        shutil.copytree(ROOT / "inputs", self.inputs, dirs_exist_ok=True)

    def tearDown(self):
        self.temporary.cleanup()

    def replace_setting(self, filename, old, new):
        path = self.inputs / filename
        content = path.read_text(encoding="utf-8")
        self.assertIn(old, content)
        path.write_text(content.replace(old, new), encoding="utf-8")

    def test_hardware_changes_reach_studies_and_verification(self):
        current = NOTEBOOK.load_config(self.inputs)
        diameter = current["vehicle"]["diameter_m"]
        startup = current["vehicle"]["stages"][1]["startup_propellant_kg"]
        self.replace_setting("vehicle.toml", f"diameter_m = {diameter}",
                             "diameter_m = 6.2")
        self.replace_setting("vehicle.toml", f"startup_propellant_kg = {startup}",
                             "startup_propellant_kg = 123.0")
        checks = NOTEBOOK.load_checks(self.inputs)
        self.assertEqual(checks["integrated_tmi"]["inputs"]["restart_propellant_kg"], 123.0)
        flight = NOTEBOOK.load_config(self.inputs)
        self.assertEqual(flight["vehicle"]["diameter_m"], 6.2)
        sizing = checks["sourced_sizing"]["inputs"]
        self.assertEqual(sizing["tanks"][0]["loaded_mass_kg"],
                         flight["vehicle"]["stages"][0]["tanks"][1]["loaded_mass_kg"])
        self.assertEqual(checks["PROP.3.TANK.02.A01"]["inputs"], sizing)
        self.assertEqual(checks["PERF.2.32.A02"]["inputs"], sizing)

    def test_choices_display_actual_changed_values(self):
        selected = NOTEBOOK.load_config(self.inputs)["vehicle"]["fairing_length_m"]
        self.replace_setting("vehicle.toml", f"fairing_length_m = {selected}",
                             "fairing_length_m = 12.25")
        choices = NOTEBOOK.choices_table(self.inputs).set_index("Choice")
        self.assertIn("12.25 m", choices.loc["Vehicle geometry", "Selection"])
        self.assertNotIn(str(selected), choices.loc["Vehicle geometry", "Selection"])

    def test_inconsistent_stage_and_tank_propellant_is_rejected(self):
        stage = NOTEBOOK.load_config(self.inputs)["vehicle"]["stages"][0]
        total = stage["propellant_kg"]
        self.replace_setting("vehicle.toml", f"propellant_kg = {total}",
                             f"propellant_kg = {total + 100}")
        with self.assertRaisesRegex(ValueError, "Stage 1: propellant_kg.*Update both"):
            NOTEBOOK.load_config(self.inputs)

    def test_propellant_source_roundoff_is_allowed(self):
        stage = NOTEBOOK.load_config(self.inputs)["vehicle"]["stages"][0]
        total = stage["propellant_kg"]
        self.replace_setting("vehicle.toml", f"propellant_kg = {total}",
                             f"propellant_kg = {total + 1e-7}")
        self.assertEqual(NOTEBOOK.load_config(self.inputs)["vehicle"]["stages"][0]["propellant_kg"],
                         total + 1e-7)

    def test_installed_thrust_checks_follow_both_selected_stages(self):
        config = NOTEBOOK.load_config(self.inputs)
        for stage, selected_thrust in zip(config["vehicle"]["stages"], (7200000, 990000)):
            self.replace_setting("vehicle.toml", f"thrust_N = {stage['thrust_N']}",
                                 f"thrust_N = {selected_thrust}")
        checks = NOTEBOOK.load_checks(self.inputs)
        for activity, expected in (("PROP.3.SYS.04.A02", 7200000),
                                   ("PROP.3.SYS.05.A02", 990000)):
            inputs = checks[activity]["inputs"]
            self.assertEqual(inputs["per_engine_thrust_N"] * inputs["engine_count"], expected)

    def test_display_covers_every_flight_setting(self):
        def leaves(value, path=""):
            if isinstance(value, dict):
                for key, item in value.items():
                    yield from leaves(item, f"{path}.{key}" if path else key)
            elif isinstance(value, list) and any(isinstance(item, dict) for item in value):
                for index, item in enumerate(value):
                    yield from leaves(item, f"{path}.{index}")
            else:
                yield path

        tables = NOTEBOOK.input_tables(self.inputs)
        displayed = {key for table in tables.values() for key in table["Input"]}
        self.assertTrue(set(leaves(NOTEBOOK.load_config(self.inputs))).issubset(displayed))
        for table in tables.values():
            self.assertTrue(table["Source"].str.contains(".toml", regex=False).all())
            self.assertFalse(any(isinstance(value, (dict, list)) for value in table["Value"]))
        self.assertIn("activities.PERF.2.06.A01.acceptance.0.limit", displayed)
        self.assertIn("studies.flight_review.max_qalpha_aoa_deg", displayed)

    def test_mission_vectors_are_reused_and_studies_are_not_evidence(self):
        checks = NOTEBOOK.load_checks(self.inputs)
        mission = NOTEBOOK.load_mission_cases(self.inputs)
        self.assertEqual(checks["lambert_window"]["inputs"]["cases"], mission["cases"])
        self.assertNotIn("studies", checks)
        self.assertNotIn("tank_thickness", checks)
        self.assertFalse(any("example" in key for key in checks))
        targets = NOTEBOOK.load_study("departure_targets", self.inputs)["targets"]
        self.assertEqual(targets[1]["target_c3_km2_s2"],
                         checks["integrated_tmi"]["inputs"]["target_c3_km2_s2"])

    def test_returns_independent_dictionaries(self):
        checks = NOTEBOOK.load_checks(self.inputs)
        original = checks["sourced_sizing"]["inputs"]["tanks"][0]["loaded_mass_kg"]
        checks["PERF.2.32.A02"]["inputs"]["tanks"][0]["loaded_mass_kg"] = -1
        self.assertEqual(checks["sourced_sizing"]["inputs"]["tanks"][0]["loaded_mass_kg"], original)
        self.assertEqual(NOTEBOOK.load_checks(self.inputs)["PERF.2.32.A02"]["inputs"]["tanks"][0]["loaded_mass_kg"], original)

    def test_explicit_json_records_still_load(self):
        record = self.inputs / "saved_run.json"
        record.write_text(json.dumps({"custom": {"input": 42}}), encoding="utf-8")
        self.assertEqual(NOTEBOOK.load_config(record), {"custom": {"input": 42}})
        self.assertEqual(NOTEBOOK.load_checks(record), {"custom": {"input": 42}})

    def test_invalid_references_fail_before_calculation(self):
        with self.assertRaisesRegex(ValueError, "Circular"):
            NOTEBOOK._resolve({"ref": "a"}, {"a": {"ref": "b"}, "b": {"ref": "a"}})
        with self.assertRaisesRegex(ValueError, "Unknown input reference"):
            NOTEBOOK._resolve({"ref": "unknown"}, {})
        with self.assertRaisesRegex(ValueError, "must be numeric"):
            NOTEBOOK._resolve({"ref": "name", "scale": 2}, {"name": "engine"})
        with self.assertRaisesRegex(ValueError, "only contain ref and scale"):
            NOTEBOOK._resolve({"ref": "x", "expression": "x + 1"}, {"x": 2})


class PhysicsChecks(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = NOTEBOOK.load_config()

    def test_body_rate_quaternion_convention(self):
        # Independent rotation-composition check catches body/world convention swaps.
        q = Rotation.from_euler("xyz", [0.3, -0.4, 0.7]).as_quat()
        w = np.array([0.2, -0.3, 0.4])
        dt = 1e-6
        advanced = NOTEBOOK.matrix(q + NOTEBOOK.quaternion_derivative(q, w) * dt)
        reference = (Rotation.from_quat(q) * Rotation.from_rotvec(w * dt)).as_matrix()
        np.testing.assert_allclose(advanced, reference, atol=2e-13, rtol=0)

    def test_initial_frame_and_earth_corotation(self):
        sim = NOTEBOOK.Simulation(self.config, ROOT)
        y = sim.initial()
        q = y[6:10]
        np.testing.assert_allclose(NOTEBOOK.matrix(q).T @ NOTEBOOK.matrix(q), np.eye(3), atol=1e-14)
        self.assertAlmostEqual(np.linalg.det(NOTEBOOK.matrix(q)), 1.0, places=13)
        np.testing.assert_allclose(
            NOTEBOOK.matrix(q)[:, 0], y[:3] / np.linalg.norm(y[:3]), atol=1e-14
        )
        # Body angular-rate vector must map to inertial Earth's spin.
        np.testing.assert_allclose(
            NOTEBOOK.matrix(q) @ y[10:13], [0, 0, 7.292115e-5], atol=1e-16
        )

    def test_standard_atmosphere_reference_points(self):
        atm = NOTEBOOK.Atmosphere(
            dict(self.config["atmosphere"], temperature_offset_K=0, density_scale=1)
        )
        T, p, rho, a, vis = atm.sample(0)
        self.assertAlmostEqual(T, 288.15, places=9)
        self.assertAlmostEqual(p, 101325, places=6)
        self.assertAlmostEqual(rho, 1.225, places=5)
        # 11 km geopotential is 11019.0678 m geometric in this reference model.
        z = 6356766 * 11000 / (6356766 - 11000)
        T, p, *_ = atm.sample(z)
        self.assertAlmostEqual(T, 216.65, places=8)
        self.assertAlmostEqual(p, 22632.04, delta=0.1)
        pressures = [atm.sample(h)[1] for h in [0, 11000, 20000, 47000, 86000, 150000]]
        self.assertTrue(all(a > b > 0 for a, b in zip(pressures, pressures[1:])))

    def test_aero_force_projection_and_domain(self):
        table = NOTEBOOK.AeroTable(ROOT / self.config["aero"]["file"], self.config["aero"])
        source = np.genfromtxt(
            ROOT / self.config["aero"]["file"], delimiter=",", names=True
        )
        for mach, alpha in [(0.5, 2.0), (1.0, 4.0), (2.0, 2.0), (10.0, 4.0)]:
            row = source[(source["mach"] == mach) & (source["alpha_deg"] == alpha)][0]
            ca, cn, cp, off = table.sample(mach, math.radians(alpha), True)
            a = math.radians(alpha)
            self.assertFalse(off)
            self.assertAlmostEqual(
                ca * math.cos(a) + cn * math.sin(a), row["cd_on"], places=13
            )
            self.assertAlmostEqual(
                cn * math.cos(a) - ca * math.sin(a), row["cl"], places=13
            )
            self.assertAlmostEqual(cp, row["cp_raw_in"] * 0.0254, places=12)
        self.assertFalse(table.sample(24.99, math.radians(4), True)[-1])
        self.assertTrue(table.sample(25.0, math.radians(4), True)[-1])
        self.assertTrue(table.sample(0.001, 0.0, True)[-1])
        self.assertTrue(table.sample(1.0, math.radians(4.01), True)[-1])

    def test_mass_closure_and_mass_shedding(self):
        c = self.config["vehicle"]
        v = NOTEBOOK.Vehicle(c)
        p = np.array([s["propellant_kg"] for s in c["stages"]])
        mass, cg, I = v.properties(p, True, True)
        expected = (
            sum(s["dry_mass_kg"] + s["propellant_kg"] for s in c["stages"])
            + c["payload_kg"]
            + c["fairing_kg"]
        )
        self.assertAlmostEqual(mass, expected, places=7)
        self.assertTrue(0 < cg < sum(s["length_m"] for s in c["stages"]))
        self.assertTrue(np.all(I > 0))
        after = v.properties(p, False, True)
        self.assertAlmostEqual(
            mass - after[0], c["stages"][0]["dry_mass_kg"] + p[0], places=7
        )
        bare = v.properties(p, False, False)
        self.assertAlmostEqual(after[0] - bare[0], c["fairing_kg"], places=7)
        p[0] -= 1000
        self.assertAlmostEqual(mass - v.properties(p, True, True)[0], 1000, places=7)

    def test_each_simulation_owns_its_configuration(self):
        config = NOTEBOOK.load_config()
        sim = NOTEBOOK.Simulation(config)
        config["guidance"]["pitch_deg"][1] = -10
        self.assertEqual(sim.c["guidance"]["pitch_deg"][1], 90)
        self.assertEqual(NOTEBOOK.load_config()["guidance"]["pitch_deg"][1], 90)


class PropulsionChecks(unittest.TestCase):
    def test_installed_thrust_uses_identical_engine_count(self):
        result = NOTEBOOK.thrust_sum(per_engine_thrust_N=2255000, engine_count=3)
        self.assertEqual(result["installed_thrust_N"], 6765000)
        self.assertEqual(result["engine_count"], 3)
        self.assertEqual(NOTEBOOK.thrust_sum(per_engine_thrust_N=0, engine_count=1)["installed_thrust_N"], 0)

    def test_engine_count_and_thrust_must_be_valid(self):
        for count in (0, -1, 1.5):
            with self.subTest(count=count), self.assertRaises(ValueError):
                NOTEBOOK.thrust_sum(per_engine_thrust_N=100, engine_count=count)
        for thrust in (-1, float("nan"), float("inf")):
            with self.subTest(thrust=thrust), self.assertRaises(ValueError):
                NOTEBOOK.thrust_sum(per_engine_thrust_N=thrust, engine_count=1)

    def test_liquid_volume_and_ullage_close_the_tank_volume(self):
        result = NOTEBOOK.liquid_ullage(tank_volume_m3=10, loaded_liquid_mass_kg=900, liquid_density_kg_m3=100)
        self.assertEqual(result["liquid_volume_m3"], 9)
        self.assertEqual(result["ullage_volume_m3"], 1)
        self.assertEqual(result["ullage_fraction"], 0.1)
        self.assertTrue(result["geometric_fill_pass"])

    def test_empty_full_and_overfilled_tanks_report_actual_inventory(self):
        for mass, expected_fraction, expected_pass in ((0, 1, True), (1000, 0, True), (1100, -0.1, False)):
            with self.subTest(mass=mass):
                result = NOTEBOOK.liquid_ullage(tank_volume_m3=10, loaded_liquid_mass_kg=mass, liquid_density_kg_m3=100)
                self.assertAlmostEqual(result["ullage_fraction"], expected_fraction)
                self.assertEqual(result["geometric_fill_pass"], expected_pass)

    def test_liquid_inputs_reject_nonphysical_or_nonfinite_values(self):
        baseline = dict(tank_volume_m3=10, loaded_liquid_mass_kg=900, liquid_density_kg_m3=100)
        for key in baseline:
            invalid_values = (-1, float("nan"), float("inf"))
            if key != "loaded_liquid_mass_kg":
                invalid_values += (0,)
            for value in invalid_values:
                with self.subTest(key=key, value=value), self.assertRaises(ValueError):
                    NOTEBOOK.liquid_ullage(**dict(baseline, **{key: value}))


class MissionChecks(unittest.TestCase):
    @staticmethod
    def cutoff_case(fairing=False):
        config = NOTEBOOK.load_config()
        sim = NOTEBOOK.Simulation(config)
        sim.phase = 2
        sim.attached = False
        sim.fairing = fairing
        y = sim.initial()
        y[:3] = [NOTEBOOK.RE + 200000, 0, 0]
        y[3:6] = [0, math.sqrt(NOTEBOOK.MU / (NOTEBOOK.RE + 200000)), 0]
        y[6:10] = Rotation.from_euler("z", 90, degrees=True).as_quat()
        y[10:13] = 0
        y[13] = 0
        y[14] = config["vehicle"]["stages"][1]["residual_kg"] + 0.05
        y[15:] = 0
        return config, sim.evaluate(0, y, True)

    def test_lambert_circular_quarter_and_long_way(self):
        radius = 7e6
        speed = math.sqrt(NOTEBOOK.MU / radius)
        period = 2 * math.pi * math.sqrt(radius**3 / NOTEBOOK.MU)
        for long_way, duration, sign in [
            (False, period / 4, 1),
            (True, period * 3 / 4, -1),
        ]:
            result = NOTEBOOK.lambert(
                departure_position_m=[radius, 0, 0],
                arrival_position_m=[0, radius, 0],
                time_of_flight_s=duration,
                central_mu_m3_s2=NOTEBOOK.MU,
                long_way=long_way,
            )
            np.testing.assert_allclose(
                result["departure_velocity_m_s"], [0, sign * speed, 0], atol=1e-5
            )
            self.assertLess(abs(result["time_of_flight_residual_s"]), 1e-6)
            json.dumps(result, allow_nan=False)

    def test_lambert_propagated_endpoint(self):
        r1 = np.array([5000, 10000, 2100]) * 1000.0
        r2 = np.array([-14600, 2500, 7000]) * 1000.0
        result = NOTEBOOK.lambert(
            departure_position_m=r1,
            arrival_position_m=r2,
            time_of_flight_s=3600,
            central_mu_m3_s2=NOTEBOOK.MU,
        )

        def rhs(t, y):
            return np.r_[y[3:], -NOTEBOOK.MU * y[:3] / np.linalg.norm(y[:3]) ** 3]

        sol = solve_ivp(
            rhs,
            (0, 3600),
            np.r_[r1, result["departure_velocity_m_s"]],
            rtol=2e-11,
            atol=1e-7,
        )
        np.testing.assert_allclose(sol.y[:3, -1], r2, atol=0.02, rtol=0)

    def test_lambert_window_compares_planet_relative_departure_energy(self):
        radius = 7e6
        speed = math.sqrt(NOTEBOOK.MU / radius)
        duration = math.pi / 2 * math.sqrt(radius**3 / NOTEBOOK.MU)
        common = dict(
            departure_position_m=[radius, 0, 0],
            arrival_position_m=[0, radius, 0],
            time_of_flight_s=duration,
            arrival_planet_velocity_m_s=[-speed, 0, 0],
        )
        result = NOTEBOOK.lambert_window(
            cases=[
                dict(common, departure_planet_velocity_m_s=[0, 0, 0]),
                dict(common, departure_planet_velocity_m_s=[0, speed, 0]),
            ],
            central_mu_m3_s2=NOTEBOOK.MU,
        )
        self.assertEqual(result["minimum_C3_case_index"], 1)
        self.assertAlmostEqual(result["cases"][0]["departure_C3_km2_s2"], NOTEBOOK.MU / radius / 1e6)
        self.assertLess(result["cases"][1]["departure_C3_km2_s2"], 1e-12)
        self.assertLess(result["cases"][0]["arrival_excess_speed_m_s"], 1e-5)


    def test_finite_tmi_stops_at_residual_and_reports_failure(self):
        config, row = self.cutoff_case()
        result = NOTEBOOK.finite_burn_tmi(
            cutoff_row=row,
            config=config,
            coast_duration_s=0,
            restart_propellant_kg=0,
            target_c3_km2_s2=9.8,
            max_burn_duration_s=10,
            output_step_s=0.1,
        )
        self.assertEqual(result["status"], "residual_propellant_limit")
        self.assertFalse(result["target_C3_reached"])
        self.assertAlmostEqual(
            result["remaining_propellant_kg"], result["required_residual_kg"], places=7
        )
        self.assertAlmostEqual(result["consumed_propellant_kg"], 0.05, places=7)
        self.assertGreater(result["history"][-2]["thrust_N"], 0)
        self.assertEqual(result["history"][-1]["thrust_N"], 0)
        self.assertEqual(result["history"][-1]["phase"], 3)
        self.assertEqual(result["history"][-1]["t_s"], result["history"][-2]["t_s"])
        json.dumps(result, allow_nan=False)

    def test_retained_fairing_cannot_disappear_at_tmi(self):
        config, row = self.cutoff_case(fairing=True)
        with self.assertRaisesRegex(ValueError, "retains the fairing"):
            NOTEBOOK.finite_burn_tmi(
                cutoff_row=row,
                config=config,
                coast_duration_s=0,
                restart_propellant_kg=0,
                target_c3_km2_s2=9.8,
                max_burn_duration_s=10,
            )

    def test_incompatible_cutoff_mass_is_rejected(self):
        config, row = self.cutoff_case()
        row["mass_kg"] += 100
        with self.assertRaisesRegex(ValueError, "mass/CG is incompatible"):
            NOTEBOOK.finite_burn_tmi(
                cutoff_row=row,
                config=config,
                coast_duration_s=0,
                restart_propellant_kg=0,
                target_c3_km2_s2=9.8,
                max_burn_duration_s=10,
            )

    def test_already_satisfied_target_does_not_start_engine_or_debit_propellant(self):
        config, row = self.cutoff_case()
        result = NOTEBOOK.finite_burn_tmi(
            cutoff_row=row,
            config=config,
            coast_duration_s=0,
            restart_propellant_kg=100,
            target_c3_km2_s2=-100,
            max_burn_duration_s=10,
        )
        self.assertEqual(result["status"], "C3_already_at_or_above_target")
        self.assertEqual(result["consumed_propellant_kg"], 0)
        self.assertEqual(result["remaining_propellant_kg"], row["prop2_kg"])
        self.assertEqual(result["elapsed_from_ascent_cutoff_s"], 0)
        self.assertTrue(all(r["thrust_N"] == 0 for r in result["history"]))
        self.assertFalse(any(e["event"] == "TMI_restart" for e in result["events"]))
        self.assertIsNone(result["powered_tracking_capacity_screen_pass"])

    def test_target_cutoff_preserves_powered_left_limit_and_engine_off_terminal(self):
        config, row = self.cutoff_case()
        initial_c3 = (
            row["inertial_speed_m_s"] ** 2 - 2 * NOTEBOOK.MU / (NOTEBOOK.RE + row["altitude_m"])
        ) / 1e6
        result = NOTEBOOK.finite_burn_tmi(
            cutoff_row=row,
            config=config,
            coast_duration_s=0,
            restart_propellant_kg=0,
            target_c3_km2_s2=initial_c3 + 1e-6,
            max_burn_duration_s=10,
        )
        self.assertEqual(result["status"], "target_C3_reached")
        self.assertGreater(result["history"][-2]["thrust_N"], 0)
        self.assertEqual(result["history"][-1]["thrust_N"], 0)
        self.assertEqual(
            result["history"][-1]["mass_kg"], result["history"][-2]["mass_kg"]
        )


class EvidenceTests(unittest.TestCase):
    def test_missing_is_not_zero_or_pass(self):
        self.assertEqual(NOTEBOOK.check_evidence({})[0], "EVIDENCE_REQUIRED")
        with self.assertRaises(ValueError):
            NOTEBOOK.acceptance(
                {"x": float("nan")},
                [
                    {
                        "metric": "x",
                        "operator": "<=",
                        "limit": 2,
                        "units": "m",
                        "source": "test",
                    }
                ],
            )

    def test_reviewed_file_integrity(self):
        with tempfile.TemporaryDirectory() as temp:
            evidence_file = Path(temp) / "record.txt"
            evidence_file.write_text("measured evidence")
            record = dict(
                file="record.txt",
                sha256=hashlib.sha256(evidence_file.read_bytes()).hexdigest(),
                reviewer="Example reviewer",
                date="2026-10-28",
                result="pass",
                configuration_id="test only",
                acceptance_criteria="x<=1 m",
                scope="test fixture",
            )
            self.assertEqual(NOTEBOOK.check_evidence(record, temp)[0], "REVIEWED_PASS")
            self.assertEqual(
                NOTEBOOK.check_evidence(record, temp, expected_configuration="new-baseline")[0],
                "STALE_EVIDENCE",
            )
            evidence_file.write_text("changed")
            self.assertEqual(NOTEBOOK.check_evidence(record, temp)[0], "INVALID_EVIDENCE")

    def test_boundary_and_source(self):
        rule = {
            "metric": "g",
            "operator": "<=",
            "limit": 6,
            "units": "g",
            "source": "test",
        }
        self.assertTrue(NOTEBOOK.acceptance({"g": 6}, [rule]))
        self.assertFalse(NOTEBOOK.acceptance({"g": 6.001}, [rule]))
        self.assertIsNone(NOTEBOOK.acceptance({"g": 1}, []))
        rule.pop("source")
        with self.assertRaises(ValueError):
            NOTEBOOK.acceptance({"g": 1}, [rule])

    def test_measurement_limits(self):
        with tempfile.TemporaryDirectory() as temp:
            measurements = Path(temp) / "measurements.csv"
            header = (
                "case_id,quantity,value,unit,lower_limit,upper_limit,criterion_source\n"
            )
            measurements.write_text(
                header
                + "A,force,10,N,0,10,requirement\nB,force,11,N,0,10,requirement\n"
            )
            frame = NOTEBOOK.measured_limits(measurements)
            self.assertEqual(frame.within_limits.tolist(), [True, False])
            measurements.write_text(header + "A,force,,N,0,10,requirement\n")
            with self.assertRaises(ValueError):
                NOTEBOOK.measured_limits(measurements)

    def test_partial_review_cannot_close_requirement(self):
        requirement = dict(
            id="TEST.1.01",
            domain="test",
            activities=["a", "b"],
            open_parameters=[],
            source_ref="test!A1",
        )
        rows = pd.DataFrame(
            [
                dict(
                    activity_id="TEST.1.01.A01",
                    requirement_id="TEST.1.01",
                    status="REVIEWED_PASS",
                )
            ]
        )
        with patch.object(
            NOTEBOOK, "load_requirements", return_value=[requirement]
        ), patch.object(NOTEBOOK, "activities", return_value=rows):
            status = NOTEBOOK.requirement_status(rows).iloc[0]
            self.assertEqual(status.status, "OPEN")
            self.assertTrue(status.assigned_reviewed_complete)
            self.assertFalse(status.complete_activity_set)
            self.assertEqual(status.expected_activities, 2)
            self.assertEqual(status.assigned_activities, 1)
            duplicate = pd.concat([rows, rows])
            self.assertEqual(NOTEBOOK.requirement_status(duplicate).iloc[0].status, "OPEN")

    def test_every_activity_has_valid_dispatch(self):
        functions, assignments = NOTEBOOK.catalog()
        self.assertEqual(assignments, {})
        self.assertEqual(set(functions), {"flight_summary", "integrated_tmi", "lambert_window", "sourced_sizing", "thrust_sum"})
        with tempfile.TemporaryDirectory() as temp:
            empty_file = Path(temp) / "empty.json"
            empty_file.write_text("{}")
            result, _ = NOTEBOOK.evaluate(
                input_file=empty_file, evidence_file=empty_file, context=None
            )
            self.assertEqual(set(result.activity_id), set(NOTEBOOK.activities().activity_id))
            self.assertFalse(result.status.str.contains("PASS").any())

    def test_missing_optional_evidence_file_leaves_activities_unreviewed(self):
        with tempfile.TemporaryDirectory() as temporary:
            folder = Path(temporary)
            empty_inputs = folder / "checks.json"
            empty_inputs.write_text("{}", encoding="utf-8")
            results, _ = NOTEBOOK.evaluate(
                input_file=empty_inputs, evidence_file=folder / "not_supplied.json",
                context=None,
            )
        self.assertEqual(set(results.activity_id), set(NOTEBOOK.activities().activity_id))
        self.assertFalse(results.status.str.startswith("REVIEWED").any())
        self.assertFalse(results.status.str.contains("PASS").any())


class ScopeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.evidence = self.folder / "evidence.json"
        self.evidence.write_text("{}", encoding="utf-8")
        self.requirements = [{
            "id": "TEST.1.01", "domain": "test", "statement": "Test requirement",
            "activities": ["Assigned first.", "Assigned second.", "Unassigned activity."],
            "methods": ["Analysis"], "source_ref": "test!G5", "open_parameters": [],
            "source_file": "test.xlsx", "source_sha256": "source digest",
        }]

    def entry(self, number=1, **changes):
        entry = {
            "activity_id": f"TEST.1.01.A{number:02d}",
            "calculation": "thrust_sum",
            "expected_activity": self.requirements[0]["activities"][number - 1],
            "artifact": "Installed thrust", "coverage": "Arithmetic only.",
        }
        return {**entry, **changes}

    def scope(self, entries):
        path = self.folder / "verification.toml"
        lines = ['description = "Test assigned scope"']
        for entry in entries:
            lines.append("\n[[activities]]")
            lines.extend(f"{key} = {json.dumps(value)}" for key, value in entry.items())
        path.write_text("\n".join(lines), encoding="utf-8")
        return path

    def test_default_scope_is_exactly_the_nineteen_assigned_activities(self):
        expected = {
            "MIS.0.01.A01", "MIS.0.01.A02", "SYS.1.02.A01", "SYS.1.08.A01",
            "SYS.1.09.A02", "SYS.1.16.A01", "PERF.2.05.A01", "PERF.2.06.A01",
            "PERF.2.13.A02", "PERF.2.32.A02", "PERF.2.35.A02",
            "PROP.3.SYS.03.A01", "PROP.3.SYS.04.A02", "PROP.3.SYS.05.A02",
            "PROP.3.SYS.16.A01", "PROP.3.TANK.01.A01", "PROP.3.TANK.02.A01",
            "PROP.3.STG1.01.A01", "PROP.3.STG2.01.A01",
        }
        table = NOTEBOOK.activities()
        self.assertEqual(set(table.activity_id), expected)
        self.assertEqual(len(table), len(expected))
        self.assertTrue(table[["artifact", "coverage"]].notna().all().all())
        self.assertEqual(set(table.kernel), set(NOTEBOOK.catalog()[0]))

    def test_scope_rejects_duplicate_unknown_route_unknown_id_and_stale_wording(self):
        invalid = [
            ([self.entry(), self.entry()], "Duplicate scoped activity"),
            ([self.entry(calculation="structures")], "Unknown calculation route"),
            ([self.entry(activity_id="TEST.1.99.A01")], "Unknown scoped activity"),
            ([self.entry(expected_activity="Assigned first")], "wording mismatch"),
        ]
        for entries, message in invalid:
            with self.subTest(message=message), patch.object(NOTEBOOK, "load_requirements", return_value=self.requirements):
                with self.assertRaisesRegex(ValueError, message):
                    NOTEBOOK.activities(self.scope(entries))

    def test_invalid_scope_stops_before_loading_inputs_or_running_calculations(self):
        scope = self.scope([self.entry(expected_activity="Old activity wording.")])
        with patch.object(NOTEBOOK, "load_requirements", return_value=self.requirements), patch.object(NOTEBOOK, "catalog") as registry, patch.object(NOTEBOOK, "load_checks") as inputs:
            with self.assertRaisesRegex(ValueError, "wording mismatch"):
                NOTEBOOK.evaluate(scope_file=scope)
        registry.assert_not_called()
        inputs.assert_not_called()

    def test_missing_assigned_result_is_distinct_from_whole_source_coverage(self):
        scope = self.scope([self.entry(), self.entry(2)])
        results = pd.DataFrame([{
            "activity_id": "TEST.1.01.A01", "requirement_id": "TEST.1.01", "status": "REVIEWED_PASS",
        }])
        with patch.object(NOTEBOOK, "load_requirements", return_value=self.requirements):
            status = NOTEBOOK.requirement_status(results, scope).iloc[0]
        self.assertEqual(status.assigned_activities, 2)
        self.assertEqual(status.expected_activities, 3)
        self.assertFalse(status.complete_assigned_set)
        self.assertFalse(status.assigned_reviewed_complete)
        self.assertFalse(status.complete_activity_set)

    def test_whole_requirement_closure_requires_reviewed_evidence_and_no_open_parameters(self):
        self.requirements[0]["activities"] = ["Assigned first."]
        scope = self.scope([self.entry()])
        results = pd.DataFrame([{
            "activity_id": "TEST.1.01.A01", "requirement_id": "TEST.1.01", "status": "REVIEWED_PASS",
        }])
        with patch.object(NOTEBOOK, "load_requirements", return_value=self.requirements):
            self.assertEqual(NOTEBOOK.requirement_status(results, scope).iloc[0].status, "REVIEWED_COMPLETE")
            results["status"] = "SCREEN_PASS"
            self.assertEqual(NOTEBOOK.requirement_status(results, scope).iloc[0].status, "OPEN")
            results["status"] = "REVIEWED_PASS"
            self.requirements[0]["open_parameters"] = ["TBD-TEST"]
            self.assertEqual(NOTEBOOK.requirement_status(results, scope).iloc[0].status, "OPEN")

    def test_only_assigned_calculations_run_and_shared_results_are_cached(self):
        scope = self.scope([self.entry(), self.entry(2)])
        calls = []

        def thrust_sum(*, per_engine_thrust_N, engine_count):
            calls.append((per_engine_thrust_N, engine_count))
            return {"thrust_N": per_engine_thrust_N * engine_count}

        def unassigned():
            self.fail("An unassigned calculation was invoked.")

        checks = {"thrust_sum": {"inputs": {"per_engine_thrust_N": 8, "engine_count": 2}}}
        with patch.object(NOTEBOOK, "load_requirements", return_value=self.requirements), patch.object(NOTEBOOK, "catalog", return_value=({"thrust_sum": thrust_sum, "flight_summary": unassigned}, {})), patch.object(NOTEBOOK, "load_checks", return_value=checks), patch.object(NOTEBOOK, "baseline_id", return_value="test"):
            result, detail = NOTEBOOK.evaluate(evidence_file=self.evidence, scope_file=scope)
            contracts = NOTEBOOK.input_contracts(scope)
        self.assertEqual(calls, [(8, 2)])
        self.assertEqual(result.activity_id.tolist(), ["TEST.1.01.A01", "TEST.1.01.A02"])
        self.assertEqual(set(result.status), {"CALCULATED"})
        self.assertEqual(set(detail), set(result.activity_id))
        self.assertEqual(contracts.calculation.tolist(), ["thrust_sum"])

    def test_shared_calculation_keeps_each_activitys_acceptance_rule(self):
        scope = self.scope([self.entry(), self.entry(2)])
        calls = []

        def calculate(*, per_engine_thrust_N, engine_count):
            calls.append((per_engine_thrust_N, engine_count))
            return {"installed_thrust_N": per_engine_thrust_N * engine_count}

        rule = {
            "metric": "installed_thrust_N", "operator": ">=", "limit": 15,
            "units": "N", "source": "Test fixture",
        }
        inputs = {"per_engine_thrust_N": 8, "engine_count": 2}
        checks = {
            "TEST.1.01.A01": {"inputs": inputs, "acceptance": [rule]},
            "TEST.1.01.A02": {
                "inputs": inputs, "acceptance": [{**rule, "limit": 17}],
            },
        }
        with patch.object(NOTEBOOK, "load_requirements", return_value=self.requirements), \
             patch.object(NOTEBOOK, "catalog", return_value=({"thrust_sum": calculate}, {})), \
             patch.object(NOTEBOOK, "load_checks", return_value=checks), \
             patch.object(NOTEBOOK, "baseline_id", return_value="test"):
            results, _ = NOTEBOOK.evaluate(
                evidence_file=self.evidence, scope_file=scope,
            )
            saved_calculation = {
                NOTEBOOK.calculation_key("thrust_sum", inputs): {"installed_thrust_N": 16},
            }
            reused_results, _ = NOTEBOOK.evaluate(
                evidence_file=self.evidence, scope_file=scope,
                calculated=saved_calculation,
            )
        self.assertEqual(calls, [(8, 2)])
        self.assertEqual(results.status.tolist(), ["SCREEN_PASS", "SCREEN_FAIL"])
        self.assertEqual(reused_results.status.tolist(), ["SCREEN_PASS", "SCREEN_FAIL"])


class RequirementImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)
        self.workbook = self.folder / "Verification.xlsx"
        self.scope = self.folder / "verification.toml"
        self.destination = self.folder / "requirements.json"
        self.scope.write_text('''description = "One assigned contribution"
[[activities]]
activity_id = "TEST.1.01.A01"
calculation = "thrust_sum"
expected_activity = "Installed thrust calculation."
artifact = "Installed thrust"
coverage = "Nominal arithmetic only."
''', encoding="utf-8")

    def write_workbook(self, duplicate_unassigned=False, stale=False):
        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)
        headers = ["ID", "Requirement", "Verification Activity", "Verification Method", "Verification Artifact"]
        for name in NOTEBOOK.SHEETS:
            sheet = workbook.create_sheet(name)
            for column, header in enumerate(headers, 1):
                sheet.cell(4, column, header)
        first = workbook[NOTEBOOK.SHEETS[0]]
        first.append(["TEST.1.01", "Selected requirement", "Changed calculation." if stale else "Installed thrust calculation.\nPhysical inspection.", "Analysis / Inspection", "Calculation and record"])
        first.append(["TEST.1.02", "Unassigned requirement TBD-LIMIT", "Unassigned analysis.", "Analysis", "Unassigned output"])
        workbook[NOTEBOOK.SHEETS[-1]].append(["TEST.1.02" if duplicate_unassigned else "TEST.1.03", "Another unassigned requirement", "Unassigned test.", "Test", "Test report"])
        workbook.save(self.workbook)
        workbook.close()

    def test_writes_selected_requirements_with_all_source_activities_and_provenance(self):
        self.write_workbook()
        summary = NOTEBOOK.import_workbook(self.workbook, self.destination, self.scope)
        records = json.loads(self.destination.read_text(encoding="utf-8"))
        self.assertEqual(summary["workbook_requirements_read"], 3)
        self.assertEqual(summary["count"], 1)
        self.assertEqual(summary["assigned_activities"], 1)
        self.assertEqual([row["id"] for row in records], ["TEST.1.01"])
        self.assertEqual(records[0]["activities"], ["Installed thrust calculation.", "Physical inspection."])
        self.assertEqual(records[0]["source_ref"], "Tier 0!C5")
        self.assertEqual(records[0]["source_file"], self.workbook.name)
        self.assertEqual(records[0]["source_sha256"], hashlib.sha256(self.workbook.read_bytes()).hexdigest())

    def test_duplicate_even_outside_assignment_rejects_entire_import(self):
        self.write_workbook(duplicate_unassigned=True)
        with self.assertRaisesRegex(ValueError, "duplicated"):
            NOTEBOOK.import_workbook(self.workbook, self.destination, self.scope)
        self.assertFalse(self.destination.exists())

    def test_stale_assignment_does_not_overwrite_previous_import(self):
        self.write_workbook(stale=True)
        self.destination.write_text("previous import", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "wording mismatch"):
            NOTEBOOK.import_workbook(self.workbook, self.destination, self.scope)
        self.assertEqual(self.destination.read_text(encoding="utf-8"), "previous import")

    def test_default_import_updates_the_bundle_without_replacing_other_inputs(self):
        self.write_workbook()
        original_workbook = self.workbook.read_bytes()
        reference_path = self.folder / "reference.json"
        reference = {
            "requirements": [], "mission_cases": {"case": "preserve"},
            "sources": {"vehicle": "preserve"},
            "reviewed_evidence": {"record": "preserve"},
            "follow_up": [{"activity": "preserve"}],
        }
        reference_path.write_text(json.dumps(reference), encoding="utf-8")
        with patch.object(NOTEBOOK, "INPUTS", self.folder):
            NOTEBOOK.import_workbook(self.workbook, scope_file=self.scope)
        updated = json.loads(reference_path.read_text(encoding="utf-8"))
        self.assertEqual([row["id"] for row in updated["requirements"]], ["TEST.1.01"])
        for key in ("mission_cases", "reviewed_evidence", "follow_up"):
            self.assertEqual(updated[key], reference[key])
        self.assertEqual(updated["sources"]["vehicle"], "preserve")
        self.assertEqual(updated["sources"]["requirements_import"]["count"], 1)
        self.assertEqual(self.workbook.read_bytes(), original_workbook)


class NotebookContractTests(unittest.TestCase):
    def test_physics_and_plots_are_visible_notebook_definitions(self):
        definitions = set()
        for cell in DOCUMENT["cells"]:
            if cell["cell_type"] != "code":
                continue
            tree = ast.parse("".join(cell["source"]))
            for node in ast.walk(tree):
                if isinstance(node, (ast.FunctionDef, ast.ClassDef)):
                    definitions.add(node.name)
                if isinstance(node, ast.ImportFrom):
                    self.assertFalse((node.module or "").startswith("model"))
                if isinstance(node, ast.Import):
                    self.assertFalse(any(item.name.startswith("model") for item in node.names))
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Name):
                    self.assertNotIn(node.func.id, {"exec", "eval", "compile"})
        self.assertTrue({
            "Atmosphere", "AeroTable", "Vehicle", "Simulation", "run_trajectory",
            "finite_burn_tmi", "lambert", "plot_results", "evaluate",
        }.issubset(definitions))

    def test_plot_results_are_displayed_in_the_notebook(self):
        code = "\n".join(
            "".join(cell["source"])
            for cell in DOCUMENT["cells"] if cell["cell_type"] == "code"
        )
        tree = ast.parse(code)
        calls = [node for node in ast.walk(tree) if isinstance(node, ast.Call)]
        self.assertTrue(any(
            isinstance(call.func, ast.Name) and call.func.id == "plot_results"
            for call in calls
        ))
        self.assertTrue(any(
            isinstance(call.func, ast.Name) and call.func.id == "display"
            for call in calls
        ))
        # An executed notebook must retain its figures for a reader on GitHub.
        executed = any(
            cell.get("execution_count") is not None
            for cell in DOCUMENT["cells"] if cell["cell_type"] == "code"
        )
        if executed:
            images = [
                output for cell in DOCUMENT["cells"]
                for output in cell.get("outputs", [])
                if "image/png" in output.get("data", {})
            ]
            self.assertGreaterEqual(len(images), 6)

    @contextmanager
    def isolated_project(self):
        """Identity tests change only a disposable copy of code and inputs."""
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            shutil.copytree(ROOT / "inputs", root / "inputs")
            notebook_path = root / "Hyperion.ipynb"
            shutil.copy2(NOTEBOOK_PATH, notebook_path)
            with patch.object(NOTEBOOK, "ROOT", root), \
                 patch.object(NOTEBOOK, "INPUTS", root / "inputs"), \
                 patch.object(NOTEBOOK, "NOTEBOOK_PATH", notebook_path):
                yield root, notebook_path

    def test_code_and_physical_input_changes_invalidate_baseline(self):
        with self.isolated_project() as (root, notebook_path):
            baseline = NOTEBOOK.baseline_id()
            document = json.loads(notebook_path.read_text(encoding="utf-8"))
            cell = next(cell for cell in document["cells"] if cell["cell_type"] == "code")
            cell["source"] = "".join(cell["source"]) + "\nTEST_SENTINEL = 1\n"
            notebook_path.write_text(json.dumps(document), encoding="utf-8")
            self.assertNotEqual(NOTEBOOK.baseline_id(), baseline)

            shutil.copy2(NOTEBOOK_PATH, notebook_path)
            self.assertEqual(NOTEBOOK.baseline_id(), baseline)
            vehicle_path = root / "inputs/vehicle.toml"
            vehicle = vehicle_path.read_text(encoding="utf-8")
            diameter = NOTEBOOK.load_config()["vehicle"]["diameter_m"]
            vehicle_path.write_text(
                vehicle.replace(f"diameter_m = {diameter}", "diameter_m = 6.2"),
                encoding="utf-8",
            )
            self.assertNotEqual(NOTEBOOK.baseline_id(), baseline)

    def test_saved_outputs_and_markdown_do_not_invalidate_baseline(self):
        with self.isolated_project() as (_, notebook_path):
            baseline = NOTEBOOK.baseline_id()
            document = json.loads(notebook_path.read_text(encoding="utf-8"))
            code_cell = next(
                cell for cell in document["cells"] if cell["cell_type"] == "code"
            )
            code_cell["execution_count"] = 7
            code_cell["outputs"] = [{
                "output_type": "stream", "name": "stdout", "text": ["Updated result\n"],
            }]
            markdown = next(
                cell for cell in document["cells"] if cell["cell_type"] == "markdown"
            )
            markdown["source"] = ["Reworded explanation; equations unchanged.\n"]
            notebook_path.write_text(json.dumps(document), encoding="utf-8")
            self.assertEqual(NOTEBOOK.baseline_id(), baseline)

    def test_reviewed_evidence_does_not_create_a_circular_baseline(self):
        with self.isolated_project() as (root, _):
            baseline = NOTEBOOK.baseline_id()
            reference_path = root / "inputs/reference.json"
            reference = json.loads(reference_path.read_text(encoding="utf-8"))
            reference["reviewed_evidence"] = {"TEST.1.01.A01": {"configuration_id": baseline}}
            reference_path.write_text(json.dumps(reference), encoding="utf-8")
            self.assertEqual(NOTEBOOK.baseline_id(), baseline)


if __name__ == "__main__":
    unittest.main()
